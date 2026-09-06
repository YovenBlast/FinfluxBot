import fitz
import pytesseract

from PIL import Image
from io import BytesIO
from docx import Document


# ==========================================
# Process PDF
# ==========================================

def process_pdf(pdf_path):

    document = fitz.open(pdf_path)

    pages = []

    total_pages = len(document)

    for page_number, page in enumerate(document):

        print(
            f"  Page {page_number + 1}/"
            f"{total_pages}"
        )

        # ======================================
        # Normal PDF text extraction
        # ======================================

        text = page.get_text("text")

        # ======================================
        # If PDF has little/no text,
        # perform OCR
        # ======================================

        if not text or len(text.strip()) < 30:

            print(
                "    No usable text found."
                " Running OCR..."
            )

            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(2, 2)
            )

            image = Image.open(
                BytesIO(
                    pixmap.tobytes("png")
                )
            )

            text = pytesseract.image_to_string(
                image
            )

        # ======================================
        # Store page
        # ======================================

        if text and text.strip():

            pages.append({
                "page": page_number + 1,
                "text": text.strip()
            })

    document.close()

    return pages


# ==========================================
# Process Word DOCX
# ==========================================

def process_docx(docx_path):

    document = Document(docx_path)

    sections = []

    current_text = []

    section_number = 1

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if not text:
            continue

        current_text.append(text)

        # --------------------------------------
        # Create a new searchable section
        # every few paragraphs
        # --------------------------------------

        if len(current_text) >= 8:

            sections.append({
                "page": section_number,
                "text": "\n".join(current_text)
            })

            current_text = []

            section_number += 1

    # ==========================================
    # Store remaining paragraphs
    # ==========================================

    if current_text:

        sections.append({
            "page": section_number,
            "text": "\n".join(current_text)
        })

    # ==========================================
    # Also extract tables
    # ==========================================

    for table_index, table in enumerate(
        document.tables,
        start=1
    ):

        table_text = []

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:

                    row_text.append(cell_text)

            if row_text:

                table_text.append(
                    " | ".join(row_text)
                )

        if table_text:

            sections.append({
                "page": f"table-{table_index}",
                "text": "\n".join(table_text)
            })

    return sections


# ==========================================
# Process Screenshot / Image
# ==========================================

def process_image(image_path):

    image = Image.open(
        image_path
    )

    # ======================================
    # Convert to RGB
    # ======================================

    if image.mode != "RGB":

        image = image.convert(
            "RGB"
        )

    # ======================================
    # OCR
    # ======================================

    text = pytesseract.image_to_string(
        image
    )

    return text.strip()